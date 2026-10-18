from io import BytesIO

import streamlit as st
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

import io
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd
import requests
import streamlit as st
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches, Pt
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from xml.sax.saxutils import escape

APP_TITLE = "Personal CMA & CA Office AI Agent"
OLLAMA_URL = "http://localhost:11434/api/generate"
OLLAMA_CHAT_URL = "http://localhost:11434/api/chat"

CHAT_SYSTEM_PROMPT = (
    "You are a careful, practical assistant for Indian income tax, GST, Tally, Zoho Books, "
    "accounting, and commercial law. Explain concepts clearly, show calculations step by step "
    "when relevant, and ask for missing facts such as financial year, GST registration type, "
    "transaction details, or software version when they affect the answer. Laws, rates, forms, "
    "deadlines, and software features can change: do not invent section numbers, notifications, "
    "rates, due dates, or legal conclusions. State uncertainty and recommend checking current "
    "official sources when an answer depends on current rules. Treat answers as general "
    "information and study/work assistance, not a substitute for advice from a qualified "
    "professional. Do not claim to know everything or to have checked live sources."
)

# -------------------- File reading --------------------
def read_pdf(data: bytes) -> str:
    text_parts = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for i, page in enumerate(pdf, start=1):
            text = page.get_text("text")
            if text.strip():
                text_parts.append(f"\n\n--- Page {i} ---\n{text}")
    return "\n".join(text_parts).strip()


def read_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table_no, table in enumerate(doc.tables, start=1):
        parts.append(f"\n[Table {table_no}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
    return "\n".join(parts).strip()


def read_xlsx(data: bytes) -> str:
    book = pd.ExcelFile(io.BytesIO(data))
    parts = []
    for sheet in book.sheet_names:
        df = pd.read_excel(book, sheet_name=sheet, header=None, dtype=str).fillna("")
        parts.append(f"\n--- Sheet: {sheet} ({len(df)} rows x {len(df.columns)} columns) ---")
        # Avoid flooding the context; retain up to 20,000 cells per sheet.
        max_rows = max(1, min(len(df), 20000 // max(1, len(df.columns))))
        parts.append(df.iloc[:max_rows].to_string(index=False, header=False))
        if max_rows < len(df):
            parts.append(f"[Note: sheet truncated for text view; first {max_rows} rows shown.]")
    return "\n".join(parts).strip()


def read_uploaded_file(uploaded) -> tuple[str, str]:
    data = uploaded.getvalue()
    name = uploaded.name
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return read_pdf(data), "PDF"
    if suffix == ".docx":
        return read_docx(data), "Word"
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx(data), "Excel"
    if suffix in (".txt", ".md", ".csv"):
        return data.decode("utf-8", errors="replace"), suffix[1:].upper()
    raise ValueError("Supported formats: PDF, DOCX, XLSX/XLSM, TXT, MD, CSV. Legacy DOC/XLS files must be saved as DOCX/XLSX first.")


# -------------------- Local AI (Ollama) --------------------
def ollama_models():
    try:
        response = requests.get("http://localhost:11434/api/tags", timeout=3)
        response.raise_for_status()
        return [m.get("name", "") for m in response.json().get("models", [])]
    except Exception:
        return []


def ask_local_ai(prompt: str, model: str, temperature: float = 0.2) -> str:
    response = requests.post(
        OLLAMA_URL,
        json={"model": model, "prompt": prompt, "stream": False,
              "options": {"temperature": temperature, "num_ctx": 8192}},
        timeout=600,
    )
    response.raise_for_status()
    return response.json().get("response", "").strip()


def ask_local_chat(messages: list[dict[str, str]], model: str) -> str:
    response = requests.post(
        OLLAMA_CHAT_URL,
        json={
            "model": model,
            "messages": [
                {"role": "system", "content": CHAT_SYSTEM_PROMPT},
                *messages[-12:],
            ],
            "stream": False,
            "options": {"temperature": 0.2, "num_ctx": 8192},
        },
        timeout=600,
    )
    response.raise_for_status()
    payload = response.json()
    message = payload.get("message") if isinstance(payload, dict) else None
    answer = message.get("content", "").strip() if isinstance(message, dict) else ""
    if not answer:
        raise ValueError("The local model returned an empty answer. Please try again.")
    return answer


def chunk_text(text: str, max_chars: int = 12000):
    text = text.strip()
    if not text:
        return []
    chunks = []
    while len(text) > max_chars:
        cut = text.rfind("\n", 0, max_chars)
        if cut < max_chars // 2:
            cut = max_chars
        chunks.append(text[:cut])
        text = text[cut:].strip()
    if text:
        chunks.append(text)
    return chunks


def run_document_task(text: str, task: str, model: str, detail: str) -> str:
    if not text.strip():
        return "No readable text was extracted. If this is a scanned PDF, run OCR first or use a searchable PDF."
    pieces = chunk_text(text)
    if len(pieces) == 1:
        prompt = make_prompt(task, detail, pieces[0])
        return ask_local_ai(prompt, model)
    partials = []
    for i, piece in enumerate(pieces, start=1):
        prompt = (
            "You are assisting a CMA finalist and staff member at a CA firm. "
            "Summarise this document segment accurately. Preserve names, dates, figures, "
            "section references and caveats. Do not invent facts. Label uncertainty.\n\n"
            f"Segment {i} of {len(pieces)}:\n{piece}"
        )
        partials.append(ask_local_ai(prompt, model))
    combined = "\n\n".join(f"Segment notes {i+1}:\n{x}" for i, x in enumerate(partials))
    final_prompt = make_prompt(task, detail, combined)
    return ask_local_ai(final_prompt, model)


def make_prompt(task: str, detail: str, text: str) -> str:
    common = (
        "You are a careful study and office assistant for a CMA finalist working in a CA firm. "
        "Use only the supplied source. Never fabricate laws, sections, case law, accounting values, "
        "tax conclusions or missing data. If information is missing or ambiguous, explicitly say so. "
        "Keep source figures unchanged and distinguish facts, assumptions, and recommendations. "
        "This is a draft for human review, not professional certification.\n\n"
    )
    instructions = {
        "Summarise": "Create a structured summary with headings, key points, important figures/dates, and a short conclusion.",
        "Analyse": "Analyse the supplied content. Separate facts, trends, inconsistencies, risks, missing information, and practical next steps. Do not claim a legal/tax conclusion without enough evidence.",
        "Study notes": "Create detailed CMA exam notes: definitions, concepts, formulas, step-by-step explanations, examples only if supported, likely exam points, and a quick revision section.",
        "CA office review": "Prepare a professional working-paper style review: purpose, source documents/data, key observations, exceptions, reconciliation questions, risks, documents to request, and reviewer checklist. Do not assert compliance from incomplete data.",
        "Action points": "Extract an actionable checklist with task, reason, source reference if available, and information still required.",
        "Translate to simple English": "Rewrite the content in clear, simple English without changing the meaning or figures."
    }
    return (
        common + f"Task: {instructions.get(task, task)}\n"
        f"Detail level: {detail}\n"
        "Output in clean Markdown. Use tables only when they improve clarity.\n\n"
        "SOURCE TEXT:\n" + text
    )


# -------------------- Export helpers --------------------
def make_docx(title: str, content: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Created: {datetime.now().strftime('%d %b %Y, %H:%M')}")
    for line in content.splitlines():
        line = line.rstrip()
        if not line:
            doc.add_paragraph("")
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", line), style="List Number")
        else:
            doc.add_paragraph(line)
    for p in doc.paragraphs:
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(10)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_pdf(title: str, content: str) -> bytes:
    out = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER,
                              textColor=colors.HexColor("#17365D"), spaceAfter=12))
    styles.add(ParagraphStyle(name="BodyCustom", parent=styles["BodyText"], leading=14, spaceAfter=6))
    story = [Paragraph(escape(title), styles["ReportTitle"]),
             Paragraph(escape(f"Created: {datetime.now().strftime('%d %b %Y, %H:%M')}"), styles["Italic"]),
             Spacer(1, 8)]
    for line in content.splitlines():
        line = line.rstrip()
        if not line:
            story.append(Spacer(1, 4))
        elif line.startswith("### "):
            story.append(Paragraph(escape(line[4:]), styles["Heading3"]))
        elif line.startswith("## "):
            story.append(Paragraph(escape(line[3:]), styles["Heading2"]))
        elif line.startswith("# "):
            story.append(Paragraph(escape(line[2:]), styles["Heading1"]))
        elif re.match(r"^\s*[-*]\s+", line):
            story.append(Paragraph("&#8226; " + escape(re.sub(r"^\s*[-*]\s+", "", line)), styles["BodyCustom"]))
        else:
            story.append(Paragraph(escape(line), styles["BodyCustom"]))
    SimpleDocTemplate(out, pagesize=A4, rightMargin=18*mm, leftMargin=18*mm,
                      topMargin=18*mm, bottomMargin=18*mm).build(story)
    return out.getvalue()


def make_excel(content: str, title: str) -> bytes:
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    rows = [{"No.": i + 1, "Content / Action Point": line} for i, line in enumerate(lines)]
    df = pd.DataFrame(rows)
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="AI Output", index=False, startrow=2)
        ws = writer.sheets["AI Output"]
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="17365D")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 110
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:B{max(3, len(rows)+3)}"
    return out.getvalue()


def simple_stats(text: str) -> dict:
    words = re.findall(r"\b[\w.,/%₹-]+\b", text)
    amounts = re.findall(r"(?:₹|Rs\.?\s*)[\s]*[\d,]+(?:\.\d{1,2})?", text, flags=re.I)
    dates = re.findall(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b", text)
    return {"Characters": len(text), "Words (approx.)": len(words),
            "Pages / sections markers": len(re.findall(r"--- Page \d+ ---", text)),
            "Currency-like values found": len(amounts), "Date-like values found": len(dates)}


# -------------------- Streamlit UI --------------------
st.set_page_config(page_title=APP_TITLE, page_icon="📚", layout="wide")
st.title("📚 Personal CMA & CA Office AI Agent")
st.caption("Local-first assistant for tax and accounting questions, documents, analysis and file exports.")

with st.sidebar:
    st.header("⚙️ Settings")
    models = ollama_models()
    if models:
        model = st.selectbox("Local AI model", models, index=0)
        st.success("Ollama is connected.")
    else:
        model = st.text_input("Ollama model name", value="qwen2.5:3b")
        st.warning("Ollama not detected. Install/start Ollama and download a model before using AI tasks.")
    task = st.selectbox("What do you want to do?",
                        ["Summarise", "Analyse", "Study notes", "CA office review", "Action points",
                         "Translate to simple English"])
    detail = st.select_slider("Detail level", options=["Short", "Normal", "Detailed", "Very detailed"], value="Detailed")
    st.markdown("---")
    st.caption("Privacy: files are processed in this app. AI text is sent to your local Ollama service, not an online AI API. Do not upload confidential client data unless your firm permits it.")

tab1, tab2, tab3, tab4 = st.tabs(
    ["💬 Tax & accounts chat", "📄 Document assistant", "🔄 Convert / extract", "📊 Quick text analysis"]
)

with tab1:
    st.markdown("### Ask about Indian tax, accounts and business software")
    st.caption(
        "Ask about income tax, GST, Tally, Zoho Books, accounting or commercial law. "
        "Answers use your local Ollama model and should be checked against current official guidance."
    )

    if st.button("Clear chat", key="clear_tax_chat"):
        st.session_state["tax_chat_messages"] = []
        st.rerun()

    chat_messages = st.session_state.setdefault("tax_chat_messages", [])
    for message in chat_messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    question = st.chat_input("Type your question")
    if question:
        pending_messages = [*chat_messages, {"role": "user", "content": question}]
        try:
            with st.chat_message("assistant"):
                with st.spinner("Thinking with your local model..."):
                    answer = ask_local_chat(pending_messages, model)
                st.markdown(answer)
            chat_messages.extend(
                [{"role": "user", "content": question}, {"role": "assistant", "content": answer}]
            )
        except requests.exceptions.ConnectionError:
            st.error("Cannot connect to Ollama. Start Ollama and make sure your selected model is installed.")
        except requests.exceptions.Timeout:
            st.error("The local model timed out. Try a smaller model or ask a shorter question.")
        except requests.exceptions.HTTPError as error:
            st.error(f"Ollama returned an error: {error}")
        except (ValueError, requests.exceptions.RequestException) as error:
            st.error(f"Could not get a chatbot response: {error}")

with tab2:
    uploaded = st.file_uploader("Upload a PDF, DOCX, XLSX, TXT, MD or CSV file",
                                type=["pdf", "docx", "xlsx", "xlsm", "txt", "md", "csv"],
                                key="main_upload")
    if uploaded:
        try:
            source_text, source_type = read_uploaded_file(uploaded)
            st.info(f"Read {len(source_text):,} characters from {uploaded.name} ({source_type}).")
            with st.expander("Preview extracted text", expanded=False):
                st.text_area("Extracted text", source_text[:25000], height=250)
            if st.button("✨ Run AI task", type="primary"):
                try:
                    with st.spinner("Working locally... Larger documents may take several minutes."):
                        result = run_document_task(source_text, task, model, detail)
                    st.session_state["last_result"] = result
                    st.session_state["last_title"] = f"{task} - {Path(uploaded.name).stem}"
                    st.session_state["last_source"] = uploaded.name
                except requests.exceptions.ConnectionError:
                    st.error("Cannot connect to Ollama. Open Ollama, then download a model (see README).")
                except requests.exceptions.Timeout:
                    st.error("The local model timed out. Try a smaller model or a shorter document.")
                except Exception as e:
                    st.error(f"AI task failed: {e}")
        except Exception as e:
            st.error(f"Could not read this file: {e}")

    if st.session_state.get("last_result"):
        st.subheader("Result")
        st.markdown(st.session_state["last_result"])
        title = st.session_state.get("last_title", "AI Report")
        content = st.session_state["last_result"]
        c1, c2, c3 = st.columns(3)
        c1.download_button("⬇️ Download Word", make_docx(title, content),
                           file_name="AI_Report.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        c2.download_button("⬇️ Download PDF", make_pdf(title, content),
                           file_name="AI_Report.pdf", mime="application/pdf")
        c3.download_button("⬇️ Download Excel", make_excel(content, title),
                           file_name="AI_Report.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.download_button("⬇️ Download Markdown / TXT", content,
                           file_name="AI_Report.md", mime="text/markdown")

with tab3:
    st.markdown("### Convert and extract")
    st.write("Extract text from PDFs, Word and Excel files, then export the extracted text to Word, PDF or Excel.")
    conv = st.file_uploader("Choose a file to extract", type=["pdf", "docx", "xlsx", "xlsm", "txt", "md", "csv"], key="convert_upload")
    if conv:
        try:
            extracted, kind = read_uploaded_file(conv)
            st.success(f"Extracted {len(extracted):,} characters.")
            st.text_area("Extracted content", extracted[:20000], height=250)
            safe_name = Path(conv.name).stem
            st.download_button("Download extracted text (.txt)", extracted,
                               file_name=f"{safe_name}_extracted.txt", mime="text/plain")
            st.download_button("Download as Word (.docx)", make_docx(f"Extracted from {conv.name}", extracted),
                               file_name=f"{safe_name}_extracted.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            st.download_button("Download as PDF", make_pdf(f"Extracted from {conv.name}", extracted),
                               file_name=f"{safe_name}_extracted.pdf", mime="application/pdf")
            st.download_button("Download as Excel", make_excel(extracted, f"Extracted from {conv.name}"),
                               file_name=f"{safe_name}_extracted.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        except Exception as e:
            st.error(f"Conversion failed: {e}")

with tab4:
    st.markdown("### Quick text analysis")
    pasted = st.text_area("Paste text or transaction narration here", height=220,
                          placeholder="Paste text, notes, report excerpts, or transaction descriptions...")
    if pasted.strip():
        stats = simple_stats(pasted)
        cols = st.columns(len(stats))
        for col, (label, value) in zip(cols, stats.items()):
            col.metric(label, f"{value:,}")
        amounts = re.findall(r"(?:₹|Rs\.?\s*)[\s]*[\d,]+(?:\.\d{1,2})?", pasted, flags=re.I)
        dates = re.findall(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b", pasted)
        st.write("**Currency-like values detected:**", ", ".join(amounts) if amounts else "None detected")
        st.write("**Date-like values detected:**", ", ".join(dates) if dates else "None detected")
        st.caption("These are pattern matches, not accounting validation or a bank reconciliation.")
        st.download_button("Download pasted text as TXT", pasted, file_name="pasted_text.txt", mime="text/plain")

st.markdown("---")
st.caption("Important: AI output may be wrong. Verify calculations, statutory references, tax positions, audit conclusions and client deliverables before use.")
